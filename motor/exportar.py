# -*- coding: utf-8 -*-
"""Exporta o resultado para Excel: Apontamentos (uma linha por correção), Notas (uma por nota) e Resumo."""
import io, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = '212E50'
SEV = {'erro': ('Erro', 'FBE3E1', '9B1C1C'), 'alerta': ('Conferir', 'FDF1D8', '8A5A00'), 'ok': ('Correta', 'E3F1E6', '1E6B35'),
       'cancelada': ('Cancelada', 'ECEEF2', '4A5468'), 'info': ('Info', 'E6EEF8', '2F5CA6')}
MOV = {'entrada': 'Entradas', 'saida': 'Saídas', 'frete': 'Frete de entrada', 'prestado': 'Serviços prestados', 'tomado': 'Serviços tomados',
       'diversos': 'Documentos diversos', 'indef': 'Sem movimento'}


def _cab(ws, cols, larg):
    ws.append(cols)
    for i, c in enumerate(cols, start=1):
        cel = ws.cell(row=ws.max_row, column=i)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor=NAVY)
        cel.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg[i - 1]
    r = ws.max_row
    ws.freeze_panes = 'A%d' % (r + 1)
    ws.auto_filter.ref = 'A%d:%s%d' % (r, get_column_letter(len(cols)), r)


def _d(v):
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', v or '')
    return '%s/%s/%s' % (m.group(3), m.group(2), m.group(1)) if m else (v or '')


def _sev(cel, sev):
    rot, bg, fg = SEV.get(sev, SEV['info'])
    cel.value = rot
    cel.fill = PatternFill('solid', fgColor=bg)
    cel.font = Font(bold=True, color=fg)


def _res(v):
    """Resolução salva: string antiga (data) ou {status, obs, em}."""
    if not v:
        return None
    if isinstance(v, str):
        return {'status': 'corrigido', 'obs': '', 'em': v}
    return v


def xlsx(res, corrigidos=None, apur_estado=None):
    corrigidos = corrigidos or {}
    from .consolidar import CAT_ROT
    wb = openpyxl.Workbook()
    emp = res.get('empresa', {})
    titulo = '%s  |  CNPJ %s  |  Competência %s  |  gerado em %s' % (emp.get('nome', ''), emp.get('cnpj', ''), res.get('competencia', ''), res.get('gerado_em', ''))

    ws = wb.active
    ws.title = 'Apontamentos'
    ws.append([titulo]); ws['A1'].font = Font(bold=True, size=12, color=NAVY)
    _cab(ws, ['Situação', 'Movimento', 'Documento', 'Número', 'Série', 'Emissão', 'CNPJ participante', 'Participante', 'Categoria', 'Apontamento', 'Detalhe',
              'Também (consequências)', 'Por que é erro', 'Na nota (XML)', 'No sistema', 'Como corrigir', 'Valor envolvido', 'Onde está na planilha', 'Chave de acesso',
              'Resolução', 'Justificativa / observação', 'Resolvido em'],
         [11, 17, 10, 10, 6, 11, 17, 30, 18, 34, 50, 40, 44, 22, 22, 40, 14, 34, 46, 12, 44, 16])
    for n in res['notas']:
        for a in n['aponts']:
            if a['sev'] not in ('erro', 'alerta'):
                continue
            r = _res(corrigidos.get(a.get('k', '')))
            ws.append(['', MOV.get(n['mov'], n['mov']), n['tipo'], n['numero'], n.get('serie', ''), _d(n.get('emissao', '')), n.get('part_cnpj', ''),
                       n.get('part_nome', ''), CAT_ROT.get(a.get('categoria'), ''), a['titulo'], a['detalhe'], '; '.join(a.get('tambem') or []), a.get('fundamento', ''),
                       a['xml'], a['sis'], a['correcao'], a['valor'] or None, n.get('ref_planilha') or n.get('pdf') or '', n.get('chave', ''),
                       {'corrigido': 'Corrigido', 'justificado': 'Justificado'}.get(r['status'], '') if r else 'Pendente', r['obs'] if r else '', r['em'] if r else ''])
            if r:
                _sev(ws.cell(row=ws.max_row, column=1), 'ok'); ws.cell(row=ws.max_row, column=1).value = 'Corrigido' if r['status'] == 'corrigido' else 'Justificado'
            else:
                _sev(ws.cell(row=ws.max_row, column=1), a['sev'])
            ws.cell(row=ws.max_row, column=17).number_format = '#,##0.00'
            for c in (11, 12, 13, 16, 21):
                ws.cell(row=ws.max_row, column=c).alignment = Alignment(wrap_text=True, vertical='top')

    ws2 = wb.create_sheet('Notas')
    _cab(ws2, ['Situação', 'Movimento', 'Documento', 'Número', 'Série', 'Emissão', 'Data lançamento', 'CNPJ participante', 'Participante',
               'Valor XML', 'Valor sistema', 'Diferença', 'Erros', 'Conferir', 'Apontamentos', 'Pontos não conferidos', 'Observação da nota', 'Planilha', 'Arquivo XML/PDF', 'Chave'],
          [11, 17, 10, 10, 6, 11, 12, 17, 30, 13, 13, 12, 7, 8, 60, 34, 40, 34, 30, 46])
    for n in res['notas']:
        vx, vs = n.get('valor_xml'), n.get('valor_sis')
        ws2.append(['', MOV.get(n['mov'], n['mov']), n['tipo'], n['numero'], n.get('serie', ''), _d(n.get('emissao', '')), _d(n.get('data_lanc', '')),
                    n.get('part_cnpj', ''), n.get('part_nome', ''), vx, vs, (vs or 0) - (vx or 0) if vx is not None and vs is not None else None,
                    n['n_erros'], n['n_alertas'], '; '.join(a['titulo'] for a in n['aponts'] if a['sev'] in ('erro', 'alerta')),
                    '; '.join('%s (%s)' % (c['rot'], c['obs']) for c in n.get('checklist', []) if c['st'] == 'nao'),
                    (_res(corrigidos.get('obs:' + n.get('k', ''))) or {}).get('obs', ''), n.get('ref_planilha', ''), n.get('arquivo', ''), n.get('chave', '')])
        _sev(ws2.cell(row=ws2.max_row, column=1), n['status'])
        for c in (10, 11, 12):
            ws2.cell(row=ws2.max_row, column=c).number_format = '#,##0.00'

    ws3 = wb.create_sheet('Resumo')
    ws3.append([titulo]); ws3['A1'].font = Font(bold=True, size=12, color=NAVY)
    _cab(ws3, ['Movimento', 'Notas', 'XML', 'Lançadas', 'Corretas', 'Com erro', 'Conferir', 'Não lançadas', 'Sem XML', 'Total XML', 'Total sistema', 'Diferença'],
         [22, 8, 8, 9, 9, 9, 9, 12, 9, 15, 15, 14])
    for m, v in res['movimentos'].items():
        ws3.append([v['nome'], v['total'], v['xml'], v['lancadas'], v['ok'], v['erro'], v['alerta'], v['nao_lancadas'], v['sem_xml'],
                    v['valor_xml'], v['valor_sis'], round(v['valor_sis'] - v['valor_xml'], 2)])
        for c in (10, 11, 12):
            ws3.cell(row=ws3.max_row, column=c).number_format = '#,##0.00'
    ws3.append([])
    ws3.append(['Tipo de apontamento', 'Notas', 'Ocorrências', 'Valor envolvido'])
    for c in range(1, 5):
        ws3.cell(row=ws3.max_row, column=c).font = Font(bold=True, color=NAVY)
    for r in res['ranking']:
        ws3.append([r['titulo'], r['notas'], r['qtd'], r['valor']])
        ws3.cell(row=ws3.max_row, column=4).number_format = '#,##0.00'

    if res.get('painel'):
        _aba_painel(wb, res, titulo)
    if res.get('apuracao', {}).get('linhas'):
        _aba_apuracao(wb, res, apur_estado or {}, titulo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _aba_apuracao(wb, res, est, titulo):
    ap = res['apuracao']
    ws = wb.create_sheet('Apuração ICMS-IPI', 0)
    ws.append([titulo]); ws['A1'].font = Font(bold=True, size=12, color=NAVY)
    ws.append(['Apuração de ICMS e IPI a partir da planilha do sistema'])
    ws.append([])
    cols = ['CFOP', 'Descrição', 'Valor contábil', 'Base ICMS', 'Valor ICMS', 'Base IPI', 'Valor IPI', 'Isentas', 'Outras']
    larg = [12, 40, 16, 16, 14, 16, 14, 14, 14]
    ws.append(cols)
    for i, c in enumerate(cols, start=1):
        cel = ws.cell(row=ws.max_row, column=i)
        cel.font = Font(bold=True, color='FFFFFF'); cel.fill = PatternFill('solid', fgColor=NAVY)
        ws.column_dimensions[get_column_letter(i)].width = larg[i - 1]
    campos = ['vc', 'bc_icms', 'icms', 'bc_ipi', 'ipi', 'isentas', 'outras']
    def linha(vals, bold=False, fill=None, indent=False):
        ws.append(vals)
        r = ws.max_row
        for c in range(3, 10):
            ws.cell(row=r, column=c).number_format = '#,##0.00'
        if bold or fill:
            for c in range(1, 10):
                if bold:
                    ws.cell(row=r, column=c).font = Font(bold=True, color=NAVY)
                if fill:
                    ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=fill)
        if indent:
            ws.cell(row=r, column=1).alignment = Alignment(indent=2)
            for c in range(1, 10):
                ws.cell(row=r, column=c).font = Font(color='5B6478', size=10)
    for grp, rot in (('entrada', 'ENTRADAS'), ('saida', 'SAÍDAS')):
        ls = [l for l in ap['linhas'] if l['grupo'] == grp]
        if not ls:
            continue
        linha([rot], bold=True, fill='E9F0F9')
        for l in ls:
            linha([l['cfop'], l['desc']] + [l[k] for k in campos], bold=True)
            for c in l['csts']:
                linha(['CST ' + c['cst'], c['desc']] + [c[k] for k in campos], indent=True)
        t = ap['totais'][grp]
        linha(['Total ' + rot.lower(), ''] + [t[k] for k in campos], bold=True, fill='EEF0F4')
    calc = est.get('calc') or {}
    if calc:
        ws.append([])
        ws.append(['AJUSTES DA APURAÇÃO']); ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color=NAVY)
        ws.append(['Tipo', 'Descrição', 'Valor', 'Código SPED (E111)'])
        for c in range(1, 5):
            ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
        for aj in calc.get('ajustes', []):
            ws.append([aj.get('tipo'), aj.get('descricao'), aj.get('valor'), aj.get('codigo')])
            ws.cell(row=ws.max_row, column=3).number_format = '#,##0.00'
        ws.append([])
        ws.append(['RESUMO DA APURAÇÃO']); ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color=NAVY)
        for imp in ('icms', 'ipi'):
            for rot, v in calc.get('resumo', {}).get(imp, []):
                ws.append([rot, '', v])
                ws.cell(row=ws.max_row, column=3).number_format = '#,##0.00'
            ws.append([])
    ws.freeze_panes = 'A5'


def _aba_painel(wb, res, titulo):
    P = res['painel']; f = P['faturamento']; im = P['impostos']
    ws = wb.create_sheet('Painel da empresa', 0)
    ws.append([titulo]); ws['A1'].font = Font(bold=True, size=12, color=NAVY)
    ws.column_dimensions['A'].width = 52; ws.column_dimensions['B'].width = 18; ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 10; ws.column_dimensions['E'].width = 10
    def sec(t):
        ws.append([]); ws.append([t]); ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color='FFFFFF')
        for c in range(1, 6):
            ws.cell(row=ws.max_row, column=c).fill = PatternFill('solid', fgColor=NAVY)
    def lin(*v):
        ws.append(list(v))
        for c in range(2, 4):
            ws.cell(row=ws.max_row, column=c).number_format = '#,##0.00'
    sec('FATURAMENTO'); lin('', 'XML', 'Sistema')
    lin('Faturamento bruto (vendas, inclui remessas 5116/5117)', f['bruto'], f['sistema_bruto'])
    lin('(-) Devoluções de venda', f['devolucoes'], f['sistema_devol'])
    lin('= Faturamento líquido', f['liquido'], round(f['sistema_bruto'] - f['sistema_devol'], 2))
    lin('Serviços prestados', f['servicos']); lin('= Receita líquida total', f['receita_liquida_total'])
    lin('Carga tributária (impostos destacados / receita bruta, %)', P['carga_efetiva'])
    sec('IMPOSTOS DO PERÍODO (destacados/apurados)')
    for k, r in (('icms_saidas', 'ICMS destacado nas vendas'), ('ipi_saidas', 'IPI destacado nas vendas'), ('st_saidas', 'ICMS-ST nas vendas'),
                 ('pis_saidas', 'PIS destacado nas vendas'), ('cofins_saidas', 'COFINS destacada nas vendas'), ('iss_prestados', 'ISS sobre serviços prestados'),
                 ('iss_retido_prestados', 'ISS retido pelos tomadores'), ('ret_sofridas', 'Retenções sofridas'), ('ret_a_recolher', 'Retenções a recolher (tomados)'),
                 ('iss_retido_tomados', 'ISS retido a recolher (tomados)')):
        lin(r, im.get(k, 0))
    for chave, t in (('clientes', 'CLIENTES QUE MAIS COMPRAM'), ('fornecedores_mp', 'FORNECEDORES DE MATÉRIA-PRIMA'), ('fornecedores', 'FORNECEDORES (TODAS AS COMPRAS)'),
                     ('finalidade', 'COMPRAS POR FINALIDADE'), ('produtos_vendidos', 'PRODUTOS MAIS VENDIDOS'), ('materias_primas', 'MATÉRIAS-PRIMAS MAIS COMPRADAS'),
                     ('vendas_uf', 'VENDAS POR DESTINO')):
        sec(t); lin('Nome', 'Valor', 'CNPJ/NCM', '%', 'Notas')
        for x in P.get(chave, []):
            lin(x['nome'], x['valor'], x.get('cnpj', ''), x.get('pct'), x.get('notas', ''))
    sn = P['simples']
    sec('FORNECEDORES DO SIMPLES NACIONAL')
    lin('Fornecedores', sn['qtd']); lin('Comprado deles', sn['valor']); lin('% das compras', sn['pct_compras']); lin('Crédito de ICMS destacado', sn['credito_destacado'])
    for x in sn['top']:
        lin(x['nome'], x['valor'], x.get('cnpj', ''), x.get('pct'), x.get('notas', ''))
