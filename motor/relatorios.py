# -*- coding: utf-8 -*-
"""Relatórios do fechamento, sempre em PDF e Excel:
- impostos: demonstrativo por imposto (ICMS, IPI, ICMS-ST, DIFAL venda a não contribuinte, ISS, retenções federais);
- movimento: entradas e saídas, consolidado ou por receita (produtos, serviços, locações);
- participantes: clientes e fornecedores.
Cada relatório é montado como {titulo, subtitulo, secoes:[{titulo, colunas, linhas, total, nota}]} e o mesmo
objeto vira Excel (openpyxl) e PDF (HTML paginado pelo PyMuPDF)."""
import io, os, re, datetime
from collections import defaultdict
from .util import brl, raiz
from .consolidar import CAT_ROT

WEB = os.path.join(os.path.dirname(__file__), '..', 'web')
NAVY = '212E50'
MOV = {'entrada': 'Entrada', 'saida': 'Saída', 'frete': 'Frete', 'prestado': 'Serviço prestado', 'tomado': 'Serviço tomado', 'diversos': 'Diversos', 'indef': ''}
RECEITA = {'produtos': 'Produtos', 'servicos': 'Serviços', 'locacoes': 'Locações'}
RE_LOC = re.compile(r'loca[cç][aã]o|aluguel|alugu|cess[aã]o de uso|arrendamento', re.I)


def _dt(d):
    return '%s/%s/%s' % (d[8:10], d[5:7], d[:4]) if d and len(d) >= 10 else (d or '')


def _cnpj(c):
    c = c or ''
    if len(c) == 14:
        return '%s.%s.%s/%s-%s' % (c[:2], c[2:5], c[5:8], c[8:12], c[12:])
    if len(c) == 11:
        return '%s.%s.%s-%s' % (c[:3], c[3:6], c[6:9], c[9:])
    return c


def _validas(res):
    """Notas que valem para relatório: sem canceladas/denegadas e sem as linhas de numeração; mais as locações da planilha."""
    out = []
    for x in (res.get('locacoes') or {}).get('linhas', []):
        out.append({'tipo': 'Locação', 'numero': x['numero'], 'serie': '', 'emissao': x['data'], 'part_nome': x['participante'], 'part_cnpj': x['cnpj'],
                    'part_uf': '', 'mov': 'loc_receita' if x['tipo'] == 'receita' else 'loc_despesa', 'valor_xml': x['valor'], 'valor_sis': x['valor'],
                    'natureza': 'normal', 'fiscal': {'fonte': 'planilha de locações', 'cfops': [], 'vc': x['valor'], 'desc_serv': x['descricao'],
                                                     'ret': {'irrf': x.get('irrf', 0)}}, 'locacao': True})
    for n in res['notas']:
        if n.get('natureza') in ('cancelada', 'numeracao', 'inutilizada'):
            continue
        if n['tipo'] == 'Documento':
            continue
        out.append(n)
    return out


def receita_de(n):
    f = n.get('fiscal') or {}
    if n.get('locacao'):
        return 'locacoes'
    if n['tipo'] == 'NFS-e':
        if RE_LOC.search(f.get('desc_serv') or '') or (f.get('servico') or '').startswith('03') or (f.get('servico') or '').startswith('3.'):
            return 'locacoes'
        return 'servicos'
    if n['tipo'] == 'CT-e':
        return 'servicos'
    cf = [c[1:] for c in f.get('cfops', [])]
    if any(c in ('908', '909') for c in cf) or RE_LOC.search(f.get('natOp') or ''):
        return 'locacoes'
    if cf and all(c == '933' for c in cf):
        return 'servicos'
    if n['mov'] in ('prestado', 'tomado'):
        return 'servicos'
    return 'produtos'


def _direcao(n):
    return 'saida' if n['mov'] in ('saida', 'prestado', 'loc_receita') else 'entrada'


def _ret_total(f):
    r = f.get('ret') or {}
    return round(sum(r.get(k, 0) or 0 for k in ('irrf', 'csll', 'inss', 'pis', 'cofins')), 2)


def _cab(res, titulo, sub=''):
    e = res.get('empresa', {})
    comp = (res.get('competencia') or '')
    return {'titulo': titulo, 'subtitulo': sub, 'empresa': e.get('nome', ''), 'cnpj': _cnpj(e.get('cnpj', '')), 'uf': e.get('uf', ''),
            'competencia': '%s/%s' % (comp[5:7], comp[:4]) if len(comp) >= 7 else comp, 'gerado': datetime.datetime.now().strftime('%d/%m/%Y %H:%M'), 'secoes': []}


# ---------------------------------------------------------------------------------------------------- 1. impostos

def impostos(res, apu=None):
    rel = _cab(res, 'Demonstrativo por imposto')
    ns = _validas(res)
    calc = (apu or {}).get('calc') or {}
    ap = res.get('apuracao') or {}
    tot = ap.get('totais') or {'entrada': {}, 'saida': {}}
    # ---- ICMS-ST, DIFAL, ISS, retenções a partir das notas
    st_said, st_ent, difal_uf, difal_sem = [], [], defaultdict(lambda: [0, 0.0, 0.0, 0.0]), []
    iss_prest, iss_tom, ret_prest, ret_tom = [], [], [], []
    for n in ns:
        f = n.get('fiscal') or {}
        part = n.get('part_nome') or ''
        if n['tipo'] in ('NF-e', 'NFC-e') and f.get('fonte') == 'xml':
            if f.get('st'):
                (st_said if n['mov'] == 'saida' else st_ent).append([_dt(n['emissao']), n['numero'], part, f.get('uf_dest') if n['mov'] == 'saida' else f.get('uf_emit'),
                                                                      ', '.join(f.get('cfops', [])), f.get('bc_st', 0), f.get('st', 0)])
            if n['mov'] == 'saida' and f.get('idDest') == '2' and f.get('indFinal') == '1' and f.get('indIEDest') == '9':
                u = difal_uf[f.get('uf_dest') or '?']
                u[0] += 1; u[1] += n['valor_xml'] or 0; u[2] += f.get('difal_dest', 0); u[3] += f.get('fcp_dest', 0)
                if not f.get('tem_grupo_difal'):
                    difal_sem.append([_dt(n['emissao']), n['numero'], part, f.get('uf_dest'), n['valor_xml']])
        if n['tipo'] == 'NFS-e':
            r = f.get('ret') or {}
            if n['mov'] == 'prestado':
                iss_prest.append([_dt(n['emissao']), n['numero'], part, f.get('servico', ''), n['valor_xml'], f.get('aliq_iss', 0), f.get('iss', 0), 'Sim' if f.get('iss_retido') else 'Não'])
                if _ret_total(f):
                    ret_prest.append([_dt(n['emissao']), n['numero'], part, n['valor_xml'], r.get('irrf', 0), r.get('csll', 0), r.get('pis', 0), r.get('cofins', 0), r.get('inss', 0), _ret_total(f)])
            else:
                if f.get('iss_retido'):
                    iss_tom.append([_dt(n['emissao']), n['numero'], part, f.get('servico', ''), n['valor_xml'], f.get('aliq_iss', 0), f.get('iss', 0)])
                if _ret_total(f):
                    ret_tom.append([_dt(n['emissao']), n['numero'], part, n['valor_xml'], r.get('irrf', 0), r.get('csll', 0), r.get('pis', 0), r.get('cofins', 0), r.get('inss', 0), _ret_total(f)])
    soma = lambda ls, i: round(sum(l[i] or 0 for l in ls), 2)
    # ---- ICMS / IPI (apuração)
    if calc:
        icms_deb, icms_cred, icms_saldo = calc.get('totDeb', 0), calc.get('totCred', 0), calc.get('saldo', 0)
        ipi_deb, ipi_cred, ipi_saldo = calc.get('ipiDeb', 0), round(calc.get('ipiCred', 0) + calc.get('ipiAnt', 0), 2), calc.get('ipiSaldo', 0)
    else:
        icms_deb, icms_cred = tot['saida'].get('icms', 0), tot['entrada'].get('icms', 0); icms_saldo = round(icms_deb - icms_cred, 2)
        ipi_deb, ipi_cred = tot['saida'].get('ipi', 0), tot['entrada'].get('ipi', 0); ipi_saldo = round(ipi_deb - ipi_cred, 2)
    iss_dev = round(sum(l[6] for l in iss_prest if l[7] == 'Não'), 2)
    iss_ret_prest = round(sum(l[6] for l in iss_prest if l[7] == 'Sim'), 2)
    sit = lambda v: ('a recolher', v) if v > 0.005 else ('saldo credor', -v)
    resumo = []
    s, v = sit(icms_saldo); resumo.append(['ICMS próprio', icms_deb, icms_cred, v, s, 'Apuração (aba 3), com os ajustes' if calc else 'Apuração da planilha, sem ajustes'])
    s, v = sit(ipi_saldo); resumo.append(['IPI', ipi_deb, ipi_cred, v, s, ''])
    resumo.append(['ICMS-ST retido nas vendas', soma(st_said, 6), 0, soma(st_said, 6), 'a recolher' if st_said else '', 'substituto tributário'])
    resumo.append(['ICMS-ST pago nas compras', 0, 0, soma(st_ent, 6), 'informativo', 'já recolhido pelo fornecedor'])
    dif_tot = round(sum(v[2] + v[3] for v in difal_uf.values()), 2)
    resumo.append(['DIFAL venda a não contribuinte (EC 87/2015)', dif_tot, 0, dif_tot, 'a recolher' if dif_tot else '', 'à UF de destino, por GNRE ou IE de substituto'])
    resumo.append(['ISS sobre serviços prestados', iss_dev, 0, iss_dev, 'a recolher' if iss_dev else '', 'mais %s retido pelos tomadores' % brl(iss_ret_prest) if iss_ret_prest else ''])
    resumo.append(['ISS retido de serviços tomados', soma(iss_tom, 6), 0, soma(iss_tom, 6), 'a recolher' if iss_tom else '', 'a empresa é responsável pela guia'])
    for nome, i, cod in (('IRRF', 4, 'DARF 1708 (serviços PJ)'), ('CSRF (CSLL + PIS + COFINS)', None, 'DARF 5952'), ('INSS retido (11%)', 8, 'EFD-Reinf R-2010 / DCTFWeb')):
        if i is None:
            sofr = round(soma(ret_prest, 5) + soma(ret_prest, 6) + soma(ret_prest, 7), 2); rec = round(soma(ret_tom, 5) + soma(ret_tom, 6) + soma(ret_tom, 7), 2)
        else:
            sofr, rec = soma(ret_prest, i), soma(ret_tom, i)
        resumo.append(['Retenção federal: ' + nome, rec, sofr, rec, 'a recolher' if rec else ('a compensar' if sofr else ''), '%s; retido de nós: %s (compensar)' % (cod, brl(sofr)) if sofr else cod])
    rel['secoes'].append({'titulo': 'Resumo por imposto', 'colunas': [('Imposto', 'txt', 34), ('Débito / a recolher', 'brl', 16), ('Crédito / a compensar', 'brl', 16),
                                                                       ('Saldo', 'brl', 16), ('Situação', 'txt', 12), ('Observação', 'txt', 40)], 'linhas': resumo,
                          'nota': 'Valores apurados e destacados nos documentos do mês. Os arquivos não trazem as guias pagas. Códigos de recolhimento são os usuais; confirme no caso concreto.'})
    # ---- ICMS por CFOP
    linhas = [[l['cfop'], l['desc'], l['vc'], l['bc_icms'], l['icms']] for l in ap.get('linhas', []) if l['icms'] or l['bc_icms'] or l['vc']]
    if linhas:
        rel['secoes'].append({'titulo': 'ICMS: apuração por CFOP (planilha do sistema)', 'colunas': [('CFOP', 'txt', 8), ('Descrição', 'txt', 40), ('Valor contábil', 'brl', 16),
                              ('Base ICMS', 'brl', 16), ('ICMS', 'brl', 14)], 'linhas': linhas, 'total': ['Total', '', None, None, None]})
    if calc.get('ajustes'):
        rel['secoes'].append({'titulo': 'ICMS: ajustes da apuração', 'colunas': [('Tipo', 'txt', 34), ('Descrição', 'txt', 52), ('Valor', 'brl', 14), ('Código E111', 'txt', 12)],
                              'linhas': [[a['tipo'], a['descricao'], a['valor'], a['codigo']] for a in calc['ajustes']]})
    if calc.get('resumo'):
        rel['secoes'].append({'titulo': 'ICMS e IPI: resumo da apuração', 'colunas': [('Linha', 'txt', 50), ('Valor', 'brl', 18)],
                              'linhas': [[r, v] for r, v in calc['resumo']['icms']] + [['', None]] + [[r, v] for r, v in calc['resumo']['ipi']]})
    linhas = [[l['cfop'], l['desc'], l['vc'], l['bc_ipi'], l['ipi']] for l in ap.get('linhas', []) if l['ipi']]
    if linhas:
        rel['secoes'].append({'titulo': 'IPI por CFOP', 'colunas': [('CFOP', 'txt', 8), ('Descrição', 'txt', 40), ('Valor contábil', 'brl', 16), ('Base IPI', 'brl', 16), ('IPI', 'brl', 14)],
                              'linhas': linhas, 'total': ['Total', '', None, None, None]})
    colst = [('Emissão', 'txt', 11), ('Nota', 'txt', 10), ('Participante', 'txt', 40), ('UF', 'txt', 5), ('CFOP', 'txt', 12), ('Base ST', 'brl', 15), ('ICMS-ST', 'brl', 14)]
    if st_said:
        rel['secoes'].append({'titulo': 'ICMS-ST retido nas vendas', 'colunas': colst, 'linhas': st_said, 'total': ['Total', '', '', '', '', None, None]})
    if st_ent:
        rel['secoes'].append({'titulo': 'ICMS-ST nas compras (pago pelo fornecedor)', 'colunas': colst, 'linhas': st_ent, 'total': ['Total', '', '', '', '', None, None]})
    if difal_uf:
        rel['secoes'].append({'titulo': 'DIFAL: vendas interestaduais a consumidor final não contribuinte', 'colunas': [('UF destino', 'txt', 10), ('Notas', 'int', 8),
                              ('Valor das vendas', 'brl', 16), ('DIFAL UF destino', 'brl', 16), ('FCP', 'brl', 12)],
                              'linhas': [[uf, v[0], round(v[1], 2), round(v[2], 2), round(v[3], 2)] for uf, v in sorted(difal_uf.items())], 'total': ['Total', None, None, None, None]})
    if difal_sem:
        rel['secoes'].append({'titulo': 'Atenção: venda a não contribuinte de outra UF sem o grupo de DIFAL no XML', 'colunas': [('Emissão', 'txt', 11), ('Nota', 'txt', 10),
                              ('Cliente', 'txt', 40), ('UF', 'txt', 5), ('Valor', 'brl', 15)], 'linhas': difal_sem,
                              'nota': 'Nessas notas o DIFAL (EC 87/2015, LC 190/2022) não foi destacado: confira se é devido à UF de destino.'})
    if iss_prest:
        rel['secoes'].append({'titulo': 'ISS: serviços prestados', 'colunas': [('Emissão', 'txt', 11), ('NFS-e', 'txt', 10), ('Tomador', 'txt', 38), ('Item LC 116', 'txt', 11),
                              ('Valor', 'brl', 15), ('Alíq. %', 'pct', 8), ('ISS', 'brl', 13), ('Retido', 'txt', 8)], 'linhas': iss_prest, 'total': ['Total', '', '', '', None, '', None, '']})
    if iss_tom:
        rel['secoes'].append({'titulo': 'ISS retido de serviços tomados (a recolher)', 'colunas': [('Emissão', 'txt', 11), ('NFS-e', 'txt', 10), ('Prestador', 'txt', 38),
                              ('Item LC 116', 'txt', 11), ('Valor', 'brl', 15), ('Alíq. %', 'pct', 8), ('ISS retido', 'brl', 13)], 'linhas': iss_tom, 'total': ['Total', '', '', '', None, '', None]})
    colr = [('Emissão', 'txt', 11), ('NFS-e', 'txt', 10), ('Participante', 'txt', 34), ('Valor', 'brl', 14), ('IRRF', 'brl', 12), ('CSLL', 'brl', 11), ('PIS', 'brl', 11),
            ('COFINS', 'brl', 11), ('INSS', 'brl', 12), ('Total', 'brl', 13)]
    tr = ['Total', '', ''] + [None] * 7
    if ret_tom:
        rel['secoes'].append({'titulo': 'Retenções federais a recolher (serviços tomados)', 'colunas': colr, 'linhas': ret_tom, 'total': tr})
    if ret_prest:
        rel['secoes'].append({'titulo': 'Retenções federais sofridas (serviços prestados, a compensar)', 'colunas': colr, 'linhas': ret_prest, 'total': tr})
    return rel


# ---------------------------------------------------------------------------------------------------- 2. entradas e saídas

def movimento(res, receita='consolidado', direcao='ambos'):
    sub = 'Consolidado' if receita == 'consolidado' else 'Receita: ' + RECEITA.get(receita, receita)
    rel = _cab(res, 'Relatório de entradas e saídas', sub + ('' if direcao == 'ambos' else ' · ' + ('Entradas' if direcao == 'entradas' else 'Saídas')))
    ns = [n for n in _validas(res) if receita == 'consolidado' or receita_de(n) == receita]
    cols = [('Emissão', 'txt', 11), ('Doc.', 'txt', 7), ('Número', 'txt', 10), ('Participante', 'txt', 34), ('CNPJ/CPF', 'txt', 18), ('UF', 'txt', 4),
            ('CFOP', 'txt', 11), ('Receita', 'txt', 10), ('Valor contábil', 'brl', 15), ('Base ICMS', 'brl', 14), ('ICMS', 'brl', 12), ('IPI', 'brl', 11),
            ('ICMS-ST', 'brl', 11), ('ISS', 'brl', 11), ('Retenções', 'brl', 11), ('No sistema', 'txt', 10)]
    resumo = defaultdict(lambda: [0, 0.0, 0.0, 0.0, 0.0])
    for dr, rot in (('entrada', 'Entradas'), ('saida', 'Saídas')):
        if direcao != 'ambos' and direcao[:-1] != dr:
            continue
        linhas = []
        for n in sorted([x for x in ns if _direcao(x) == dr], key=lambda x: (x.get('emissao') or '', x.get('numero') or '')):
            f = n.get('fiscal') or {}
            rc = receita_de(n)
            sis = 'Planilha' if n.get('locacao') else ('Sim' if n.get('valor_sis') is not None and n['tipo'] != 'Lançamento' else ('Sem XML' if n['tipo'] == 'Lançamento' else 'Não'))
            vc = n['valor_xml'] if n.get('valor_xml') is not None else n.get('valor_sis') or 0
            linhas.append([_dt(n.get('emissao')), n['tipo'], n['numero'], n.get('part_nome') or '', _cnpj(n.get('part_cnpj')), n.get('part_uf') or '',
                           ', '.join(f.get('cfops', [])[:3]), RECEITA[rc], vc, f.get('bc', 0), f.get('icms', 0), f.get('ipi', 0), f.get('st', 0),
                           f.get('iss', 0) if n['tipo'] in ('NFS-e', 'Lançamento') else 0, _ret_total(f), sis])
            r = resumo[(rot, RECEITA[rc])]
            r[0] += 1; r[1] += vc; r[2] += f.get('icms', 0) or 0; r[3] += f.get('ipi', 0) or 0; r[4] += f.get('iss', 0) or 0
        rel['secoes'].append({'titulo': rot, 'colunas': cols, 'linhas': linhas, 'total': ['Total', '', '', '', '', '', '', '', None, None, None, None, None, None, None, ''],
                              'nota': '' if linhas else 'Nenhum documento.'})
    rel['secoes'].insert(0, {'titulo': 'Resumo', 'colunas': [('Movimento', 'txt', 12), ('Receita', 'txt', 12), ('Documentos', 'int', 11), ('Valor contábil', 'brl', 17),
                             ('ICMS', 'brl', 14), ('IPI', 'brl', 13), ('ISS', 'brl', 13)],
                             'linhas': [[k[0], k[1], v[0], round(v[1], 2), round(v[2], 2), round(v[3], 2), round(v[4], 2)] for k, v in sorted(resumo.items())],
                             'nota': 'Sem notas canceladas ou denegadas. Valores do XML; lançamentos sem XML entram pelo valor da planilha ("Sem XML").'})
    return rel


# ---------------------------------------------------------------------------------------------------- 3. clientes e fornecedores

def participantes(res, quem='ambos', receita='consolidado'):
    rel = _cab(res, 'Relatório de clientes e fornecedores', 'Consolidado' if receita == 'consolidado' else 'Receita: ' + RECEITA.get(receita, receita))
    ns = [n for n in _validas(res) if receita == 'consolidado' or receita_de(n) == receita]
    for lado, rot in (('saida', 'Clientes'), ('entrada', 'Fornecedores')):
        if quem not in ('ambos', 'clientes' if lado == 'saida' else 'fornecedores'):
            continue
        g = {}
        for n in ns:
            if _direcao(n) != lado:
                continue
            f = n.get('fiscal') or {}
            c = n.get('part_cnpj') or ''
            k = raiz(c) if len(c) == 14 else (c or n.get('part_nome'))
            e = g.setdefault(k, {'nome': n.get('part_nome') or '', 'cnpj': c, 'uf': n.get('part_uf') or '', 'regime': '', 'notas': 0, 'valor': 0.0, 'prod': 0.0,
                                 'serv': 0.0, 'icms': 0.0, 'ipi': 0.0, 'ret': 0.0, 'estab': set()})
            if len(e['estab']) == 0 or (c and c not in e['estab']):
                e['estab'].add(c)
            if f.get('crt') in ('1', '4'):
                e['regime'] = 'Simples'
            elif f.get('crt') == '3' and not e['regime']:
                e['regime'] = 'Normal'
            vc = n['valor_xml'] if n.get('valor_xml') is not None else n.get('valor_sis') or 0
            e['notas'] += 1; e['valor'] += vc
            if receita_de(n) == 'produtos':
                e['prod'] += vc
            else:
                e['serv'] += vc
            e['icms'] += f.get('icms', 0) or 0; e['ipi'] += f.get('ipi', 0) or 0; e['ret'] += _ret_total(f)
        tot = sum(e['valor'] for e in g.values()) or 1
        linhas = [[i, e['nome'], _cnpj(e['cnpj']) + (' (+%d)' % (len(e['estab']) - 1) if len(e['estab']) > 1 else ''), e['uf'], e['regime'], e['notas'],
                   round(e['prod'], 2), round(e['serv'], 2), round(e['valor'], 2), round(100 * e['valor'] / tot, 2), round(e['icms'], 2), round(e['ipi'], 2), round(e['ret'], 2)]
                  for i, e in enumerate(sorted(g.values(), key=lambda x: -x['valor']), start=1)]
        cols = [('#', 'int', 5), ('Nome', 'txt', 38), ('CNPJ/CPF', 'txt', 22), ('UF', 'txt', 4), ('Regime', 'txt', 9), ('Notas', 'int', 7),
                ('Produtos', 'brl', 15), ('Serviços/locações', 'brl', 15), ('Total', 'brl', 15), ('% do total', 'pct', 9), ('ICMS', 'brl', 13), ('IPI', 'brl', 12), ('Retenções', 'brl', 12)]
        conc = sum(l[9] for l in linhas[:5])
        rel['secoes'].append({'titulo': '%s (%d)' % (rot, len(linhas)), 'colunas': cols, 'linhas': linhas,
                              'total': ['', 'Total', '', '', '', None, None, None, None, None, None, None, None],
                              'nota': 'Os 5 maiores somam %s%% do total. Matriz e filiais agrupadas pela raiz do CNPJ. Regime pelo CRT do XML (fornecedores).' % ('%.1f' % conc).replace('.', ',')})
    return rel


# ---------------------------------------------------------------------------------------------------- saídas: Excel e PDF

def _fmt(v, tipo):
    if v is None or v == '':
        return ''
    if tipo == 'brl':
        return '{:,.2f}'.format(v).replace(',', 'X').replace('.', ',').replace('X', '.')
    if tipo == 'pct':
        return ('%.2f' % v).replace('.', ',')
    return str(v)


def _totais(sec):
    t = sec.get('total')
    if not t:
        return None
    out = list(t)
    for i, v in enumerate(t):
        if v is None:
            out[i] = round(sum((l[i] or 0) for l in sec['linhas'] if isinstance(l[i], (int, float))), 2)
    return out


def xlsx(rel):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\\/*?:\[\]]', '', rel['titulo'])[:31]
    ws.append([rel['titulo'] + (' · ' + rel['subtitulo'] if rel['subtitulo'] else '')]); ws['A1'].font = Font(bold=True, size=13, color=NAVY)
    ws.append(['%s · CNPJ %s · Competência %s · gerado em %s' % (rel['empresa'], rel['cnpj'], rel['competencia'], rel['gerado'])])
    ws['A2'].font = Font(color='5B6478')
    larg = defaultdict(int)
    for sec in rel['secoes']:
        ws.append([])
        ws.append([sec['titulo']]); ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11, color=NAVY)
        ws.append([c[0] for c in sec['colunas']])
        for i, c in enumerate(sec['colunas'], start=1):
            cel = ws.cell(row=ws.max_row, column=i)
            cel.font = Font(bold=True, color='FFFFFF'); cel.fill = PatternFill('solid', fgColor=NAVY)
            cel.alignment = Alignment(horizontal='right' if c[1] in ('brl', 'pct', 'int') else 'left')
            larg[i] = max(larg[i], c[2])
        for l in sec['linhas']:
            ws.append(l)
            for i, c in enumerate(sec['colunas'], start=1):
                if c[1] == 'brl':
                    ws.cell(row=ws.max_row, column=i).number_format = '#,##0.00'
                elif c[1] == 'pct':
                    ws.cell(row=ws.max_row, column=i).number_format = '0.00'
        t = _totais(sec)
        if t:
            ws.append(t)
            for i, c in enumerate(sec['colunas'], start=1):
                cel = ws.cell(row=ws.max_row, column=i)
                cel.font = Font(bold=True, color=NAVY); cel.fill = PatternFill('solid', fgColor='EEF0F4')
                if c[1] == 'brl':
                    cel.number_format = '#,##0.00'
        if sec.get('nota'):
            ws.append([sec['nota']]); ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color='5B6478', size=9)
    for i, w in larg.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _esc(t):
    return str(t).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def pdf(rel):
    import fitz
    css = ('body{font-family:sans-serif;font-size:7.6pt;color:#1C2333}'
           'h2{font-size:10pt;color:#212E50;margin:10px 0 4px 0}'
           'table{border-collapse:collapse;width:100%}'
           'th{background-color:#212E50;color:#FFFFFF;font-size:7pt;padding:3px 4px;text-align:left}'
           'td{padding:2.5px 4px;border-bottom:0.5px solid #D5DAE3}'
           'td.n,th.n{text-align:right}'
           'tr.t td{font-weight:bold;background-color:#EEF0F4;color:#212E50;border-top:0.8px solid #212E50}'
           'p.nota{color:#5B6478;font-size:6.8pt;margin:3px 0 0 0}')
    partes = []
    for sec in rel['secoes']:
        h = ['<h2>%s</h2><table><tr>' % _esc(sec['titulo'])]
        for c in sec['colunas']:
            h.append('<th class="%s">%s</th>' % ('n' if c[1] in ('brl', 'pct', 'int') else '', _esc(c[0])))
        h.append('</tr>')
        for l in sec['linhas']:
            h.append('<tr>' + ''.join('<td class="%s">%s</td>' % ('n' if c[1] in ('brl', 'pct', 'int') else '', _esc(_fmt(v, c[1]))) for v, c in zip(l, sec['colunas'])) + '</tr>')
        t = _totais(sec)
        if t:
            h.append('<tr class="t">' + ''.join('<td class="%s">%s</td>' % ('n' if c[1] in ('brl', 'pct', 'int') else '', _esc(_fmt(v, c[1]))) for v, c in zip(t, sec['colunas'])) + '</tr>')
        h.append('</table>')
        if sec.get('nota'):
            h.append('<p class="nota">%s</p>' % _esc(sec['nota']))
        partes.append(''.join(h))
    html = '<body>%s</body>' % ''.join(partes)
    story = fitz.Story(html, user_css=css)
    largo = max((len(s['colunas']) for s in rel['secoes']), default=6) > 8
    mb = fitz.paper_rect('a4-l' if largo else 'a4')
    area = mb + (28, 70, -28, -34)
    buf = io.BytesIO()
    w = fitz.DocumentWriter(buf)
    mais = 1
    while mais:
        dev = w.begin_page(mb)
        mais, _ = story.place(area)
        story.draw(dev)
        w.end_page()
    w.close()
    doc = fitz.open('pdf', buf.getvalue())
    logo = os.path.join(WEB, 'logo-colorida.png')
    total = len(doc)
    for i, pg in enumerate(doc, start=1):
        r = pg.rect
        if os.path.exists(logo):
            pg.insert_image(fitz.Rect(28, 18, 128, 47), filename=logo, keep_proportion=True)
        pg.insert_text((r.width - 28 - 330, 30), rel['titulo'] + (' · ' + rel['subtitulo'] if rel['subtitulo'] else ''), fontsize=10, color=(0.13, 0.18, 0.31), fontname='hebo')
        pg.insert_text((r.width - 28 - 330, 43), '%s · CNPJ %s' % (rel['empresa'][:48], rel['cnpj']), fontsize=7.5, color=(0.36, 0.39, 0.47))
        pg.insert_text((r.width - 28 - 330, 54), 'Competência %s · gerado em %s' % (rel['competencia'], rel['gerado']), fontsize=7.5, color=(0.36, 0.39, 0.47))
        pg.draw_line((28, 60), (r.width - 28, 60), color=(0.84, 0.85, 0.89), width=0.6)
        pg.insert_text((28, r.height - 18), 'Marçal Contabilidade · Revisor do Fechamento', fontsize=7, color=(0.54, 0.58, 0.65))
        pg.insert_text((r.width - 90, r.height - 18), 'Página %d de %d' % (i, total), fontsize=7, color=(0.54, 0.58, 0.65))
    out = doc.tobytes(deflate=True, garbage=3)
    return out


def estrategico(res):
    e = res['estrategico']
    rel = _cab(res, 'Visão estratégica do fechamento')
    ev = e['evolucao']
    rel['secoes'].append({'titulo': 'Evolução mês a mês', 'colunas': [('Competência', 'txt', 12), ('Produtos (líquido)', 'brl', 17), ('Serviços', 'brl', 15), ('Locações', 'brl', 14),
                          ('Receita', 'brl', 17), ('Compras', 'brl', 16), ('Margem bruta est.', 'brl', 16), ('ICMS (+ pagar / − credor)', 'brl', 18), ('IPI (+ / −)', 'brl', 14),
                          ('Carga %', 'pct', 9), ('Notas com erro', 'int', 10)],
                          'linhas': [['%s/%s' % (x['competencia'][5:7], x['competencia'][:4]), x['produtos'], x['servicos'], x['locacoes'], x['receita'], x['compras'], x['margem'],
                                      x['icms_saldo'], x['ipi_saldo'], x['carga'], x['com_erro']] for x in ev],
                          'nota': 'Cada fechamento guardado da empresa entra aqui; o mês aberto é o último.'})
    m = e['margem']
    rel['secoes'].append({'titulo': 'Margem bruta estimada (produtos)', 'colunas': [('Linha', 'txt', 50), ('Valor', 'brl', 18)],
                          'linhas': [['Vendas líquidas de produtos', m['vendas_liquidas']], ['(−) Compras de matéria-prima e mercadoria', m['compras_mp_revenda']],
                                     ['= Margem bruta estimada', m['margem']], ['Margem %', m['margem_pct']]], 'nota': m['premissa']})
    sd = e['saldo']
    linhas = [['ICMS', 'saldo credor' if sd['icms']['credor'] else 'a recolher', sd['icms']['valor']], ['IPI', 'saldo credor' if sd['ipi']['credor'] else 'a recolher', sd['ipi']['valor']]]
    rel['secoes'].append({'titulo': 'Saldo credor', 'colunas': [('Imposto', 'txt', 10), ('Situação', 'txt', 14), ('Valor', 'brl', 16)], 'linhas': linhas,
                          'nota': (' '.join('%s: %s (%s; %s).' % (a['imposto'], a['acao'], a['base'], a['quando']) for a in sd['acoes'])) or 'Sem saldo credor no mês.'})
    rel['secoes'].append({'titulo': 'Créditos não aproveitados', 'colunas': [('Tipo', 'txt', 52), ('Notas', 'int', 8), ('Valor', 'brl', 16)],
                          'linhas': [[c['tipo'], c['notas'], c['valor']] for c in e['creditos']], 'total': ['Total', None, None],
                          'nota': 'Valores dos apontamentos ainda abertos e do crédito do Simples fora do livro.' if e['creditos'] else 'Nenhum crédito não aproveitado identificado.'})
    lr = e['lucro_real']
    rel['secoes'].append({'titulo': 'Projeção do Lucro Real (estimativa)', 'colunas': [('Linha', 'txt', 56), ('Valor', 'brl', 18)],
                          'linhas': [['Receita líquida do mês (base)', lr['receita_base']], ['PIS/COFINS: débito (9,25%)', lr['pc_debito']], ['PIS/COFINS: créditos (9,25%)', lr['pc_credito']],
                                     ['= PIS/COFINS a recolher (estimado)', lr['pc_saldo']], ['PIS/COFINS destacado nas NF-e de venda (conferência)', lr['pc_destacado_xml']],
                                     ['IRPJ: base presumida da estimativa', lr['irpj_base']], ['IRPJ estimado (15% + adicional 10%)', lr['irpj']],
                                     ['CSLL: base presumida da estimativa', lr['csll_base']], ['CSLL estimada (9%)', lr['csll']]], 'nota': ' '.join(lr['premissas'])})
    rf = e['reforma']
    cbs, ibs = rf['ref_cbs'], rf['ref_ibs']
    base = rf['receita_base']
    deb = base * cbs / 100; cr = rf['compras_credito'] * cbs / 100
    rel['secoes'].append({'titulo': 'Reforma tributária: simulação de 2027 (CBS no lugar de PIS/COFINS)', 'colunas': [('Linha', 'txt', 56), ('Valor', 'brl', 18)],
                          'linhas': [['PIS/COFINS hoje (estimado)', rf['pis_cofins_hoje']], ['CBS: débito (%s%% sobre a receita)' % ('%g' % cbs).replace('.', ','), round(deb, 2)],
                                     ['CBS: crédito das compras do regime regular', round(cr, 2)], ['= CBS a recolher (estimada)', round(deb - cr, 2)],
                                     ['Diferença CBS − PIS/COFINS', round(deb - cr - rf['pis_cofins_hoje'], 2)], ['IBS 2027 (0,1%%) sobre a receita', round(base * 0.001, 2)],
                                     ['IBS na alíquota cheia de referência (%s%%), para quando ICMS e ISS saírem' % ('%g' % ibs).replace('.', ','), round(base * ibs / 100, 2)]],
                          'nota': ' '.join(rf['premissas'])})
    rk = e['risco']
    rel['secoes'].append({'titulo': 'Concentração e risco', 'colunas': [('Indicador', 'txt', 50), ('Valor', 'txt', 40)],
                          'linhas': [['Maior cliente', ('%s (%s%%)' % (rk['top1_cliente']['nome'], ('%.1f' % rk['top1_cliente']['pct']).replace('.', ','))) if rk['top1_cliente'] else '—'],
                                     ['5 maiores clientes', ('%.1f%%' % rk['top5_clientes_pct']).replace('.', ',')],
                                     ['Maior fornecedor', ('%s (%s%%)' % (rk['top1_fornecedor']['nome'], ('%.1f' % rk['top1_fornecedor']['pct']).replace('.', ','))) if rk['top1_fornecedor'] else '—'],
                                     ['Compras de fornecedores do Simples', ('%.1f%%' % rk['simples_pct_compras']).replace('.', ',')]] + [['Atenção', a] for a in rk['alertas']]})
    op = e['operacao']
    rel['secoes'].append({'titulo': 'Indicadores do fechamento', 'colunas': [('Indicador', 'txt', 50), ('Valor', 'txt', 30)],
                          'linhas': [['Documentos (XML)', op['documentos_xml']], ['Lançamentos na planilha', op['lancamentos']], ['Notas com erro', '%s (%s%%)' % (op['notas_com_erro'], str(op['taxa_erro']).replace('.', ','))],
                                     ['Apontamentos', op['apontamentos']], ['Primeiro arquivo recebido', op['recebido_em'] or '—'], ['Calculado em', op['calculado_em']]] +
                                    [['Apontamentos: ' + CAT_ROT.get(k, k), v] for k, v in sorted(op['por_categoria'].items(), key=lambda kv: -kv[1])]})
    return rel


def estrategia_scp(res):
    k = res['scp_calc']
    c = k['cfg']
    rel = _cab(res, 'Estratégia tributária Matriz × SCP', 'Matriz (Lucro Real) e SCP %s(Lucro Presumido)' % (('CNPJ %s ' % _cnpj(c.get('scp_cnpj'))) if c.get('scp_cnpj') else ''))
    rs_, d = k['resumo'], k['resumo']['distribuicao']
    cm = k['comparativo']
    rel['secoes'].append({'titulo': 'Resumo da operação', 'colunas': [('Indicador', 'txt', 46), ('Matriz', 'brl', 16), ('SCP', 'brl', 16), ('Consolidado', 'brl', 16)],
                          'linhas': [['Receita', rs_['matriz']['receita'], rs_['scp']['receita'], rs_['matriz']['receita'] + rs_['scp']['receita']],
                                     ['PIS/COFINS a recolher', rs_['matriz']['pis_cofins'], rs_['scp']['pis_cofins'], rs_['matriz']['pis_cofins'] + rs_['scp']['pis_cofins']],
                                     ['IRPJ/CSLL (Matriz no mês; SCP no trimestre)', rs_['matriz']['irpj_csll'], rs_['scp']['irpj_csll_tri'], rs_['matriz']['irpj_csll'] + rs_['scp']['irpj_csll_tri']],
                                     ['Saldo credor de PIS/COFINS a transportar', rs_['matriz']['saldo_credor'], 0, rs_['matriz']['saldo_credor']]]})
    rel['secoes'].append({'titulo': 'Distribuição do faturamento de serviço e locação', 'colunas': [('Receita', 'txt', 20), ('Total', 'brl', 16), ('Matriz', 'brl', 16), ('% Matriz', 'pct', 10),
                          ('SCP', 'brl', 16), ('% SCP', 'pct', 10)],
                          'linhas': [['Serviços', d['serv_total'], d['serv_matriz'], d['pct_serv_matriz'], d['serv_scp'], 100 - d['pct_serv_matriz'] if d['serv_total'] else 0],
                                     ['Locação', d['loc_total'], d['loc_matriz'], d['pct_loc_matriz'], d['loc_scp'], 100 - d['pct_loc_matriz'] if d['loc_total'] else 0]],
                          'nota': ('Modo automático: a Matriz fica com o serviço/locação só até os créditos zerarem o PIS/COFINS (folga de base %s); o restante vai para a SCP. Produtos ficam sempre na Matriz.' % brl(d['folga_base']))
                          if d['modo'] == 'auto' else 'Distribuição informada manualmente. Produtos ficam sempre na Matriz.'})
    colm = [('Linha', 'txt', 52), ('Valor', 'brl', 16), ('Memória de cálculo', 'txt', 60)]
    for sec, tit in (('produtos', 'Matriz: base de débito de produtos (tratamento do ICMS)'), ('creditos', 'Matriz: base de créditos das entradas'),
                     ('matriz_pis', 'Matriz: apuração do PIS (não cumulativo)'), ('matriz_cofins', 'Matriz: apuração da COFINS (não cumulativa)'),
                     ('matriz_ir', 'Matriz: IRPJ e CSLL (Lucro Real)'), ('scp_pc', 'SCP: PIS e COFINS (cumulativo)'), ('scp_ir', 'SCP: IRPJ e CSLL (presumido, trimestral)')):
        rel['secoes'].append({'titulo': tit, 'colunas': colm, 'linhas': [[l['rot'], l['valor'], l['formula']] for l in k['memoria'].get(sec, [])]})
    rel['secoes'].append({'titulo': 'Matriz: CFOP de venda e devolução na base de débito', 'colunas': [('CFOP', 'txt', 8), ('Descrição', 'txt', 40), ('Valor contábil', 'brl', 15),
                          ('ICMS', 'brl', 13), ('IPI', 'brl', 12), ('ICMS-ST', 'brl', 12), ('Base', 'brl', 15)],
                          'linhas': [[x['cfop'], x['desc'], x['vc'], x['icms'], x['ipi'], x['st'], x['base']] for x in k['linhas_debito']], 'total': ['Total', '', None, None, None, None, None]})
    rel['secoes'].append({'titulo': 'Matriz: CFOP de entrada na base de crédito', 'colunas': [('CFOP', 'txt', 8), ('Descrição', 'txt', 38), ('Grupo', 'txt', 22), ('Valor contábil', 'brl', 15),
                          ('ICMS', 'brl', 13), ('IPI', 'brl', 12), ('Base', 'brl', 15)],
                          'linhas': [[x['cfop'], x['desc'], x['grupo'], x['vc'], x['icms'], x['ipi'], x['base']] for x in k['linhas_credito']], 'total': ['Total', '', '', None, None, None, None]})
    if k['linhas_fora']:
        rel['secoes'].append({'titulo': 'Entradas fora do crédito', 'colunas': [('CFOP', 'txt', 8), ('Descrição', 'txt', 44), ('Base', 'brl', 15)],
                              'linhas': [[x['cfop'], x['desc'], x['base']] for x in k['linhas_fora']]})
    rel['secoes'].append({'titulo': 'Comparativo tributário do mês', 'colunas': [('Cenário', 'txt', 40), ('PIS/COFINS', 'brl', 15), ('IRPJ/CSLL', 'brl', 15), ('Total', 'brl', 15), ('Carga %', 'pct', 9)],
                          'linhas': [[n, x['pis_cofins'], x['irpj_csll'], x['total'], round(100 * x['total'] / cm['receita'], 2) if cm['receita'] else 0]
                                     for n, x in (('Estratégia (Matriz + SCP)', cm['estrategia']), ('Tudo na Matriz (Lucro Real)', cm['lucro_real']), ('Tudo no Lucro Presumido', cm['presumido']))],
                          'nota': 'Economia da estratégia: %s contra tudo no Lucro Real e %s contra tudo no Presumido. %s' % (brl(cm['economia_vs_real']), brl(cm['economia_vs_presumido']), ' '.join(cm['premissas']))})
    rel['secoes'].append({'titulo': 'Alertas, riscos e premissas', 'colunas': [('Tipo', 'txt', 12), ('Descrição', 'txt', 120)],
                          'linhas': [[{'risco': 'Risco', 'alerta': 'Validar', 'info': 'Informação'}.get(a['sev'], a['sev']), a['txt']] for a in k['alertas']]})
    return rel


def gerar(res, apu, tipo, formato, **op):
    if tipo == 'scp' and formato in ('xlsx', 'pdf'):
        from . import scp_planilha
        abas = scp_planilha.montar(res, res['scp_calc'], apu)
        if formato == 'xlsx':
            return scp_planilha.xlsx(abas)
        return scp_planilha.pdf(abas, _cab(res, 'Estratégia tributária Matriz × SCP'))
    if tipo == 'scp':
        rel = estrategia_scp(res)
    elif tipo == 'estrategico':
        rel = estrategico(res)
    elif tipo == 'impostos':
        rel = impostos(res, apu)
    elif tipo == 'movimento':
        rel = movimento(res, op.get('receita', 'consolidado'), op.get('direcao', 'ambos'))
    else:
        rel = participantes(res, op.get('quem', 'ambos'), op.get('receita', 'consolidado'))
    if formato == 'json':
        for sec in rel['secoes']:
            sec['totais'] = _totais(sec)
        return rel
    return xlsx(rel) if formato == 'xlsx' else pdf(rel)
