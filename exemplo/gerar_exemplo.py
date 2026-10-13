# -*- coding: utf-8 -*-
"""Gera uma massa de teste sintética (empresa fictícia) com erros plantados de propósito em cada regra.
Uso: python3 exemplo/gerar_exemplo.py  -> cria exemplo/lote/<card>/..."""
import os, shutil, random
import openpyxl

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lote')
EMP = {'cnpj': '11222333000181', 'nome': 'COMERCIAL EXEMPLO LTDA', 'uf': 'PR', 'mun': '4106902'}
F = {
    'F1': ('22333444000155', 'DISTRIBUIDORA ALFA LTDA', 'PR', '3'),
    'F2': ('33444555000166', 'INDUSTRIA BETA SA', 'SP', '3'),
    'F3': ('44555666000177', 'GAMA COMERCIO ME', 'PR', '1'),
    'F4': ('55666777000188', 'DELTA BEBIDAS LTDA', 'SC', '3'),
    'F5': ('66777888000144', 'EPSILON AUTOPECAS LTDA', 'PR', '3'),
    'F6': ('13131313000113', 'AUTO POSTO ESTRADA LTDA', 'SP', '3'),
    'F7': ('14141414000114', 'MAQUINAS PESADAS SA', 'SP', '3'),
    'F8': ('15151515000115', 'RESTAURANTE BOA MESA LTDA', 'SP', '3'),
}
C = {'C1': ('77888999000110', 'CLIENTE UM LTDA', 'PR'), 'C2': ('88999000000120', 'CLIENTE DOIS SA', 'RS'), 'CPF': ('12345678909', 'CONSUMIDOR FINAL', 'PR')}
T = {'T1': ('99000111000130', 'TRANSPORTES RAPIDO LTDA', 'PR'), 'T2': ('10111222000140', 'LOGISTICA SUL LTDA', 'SP')}
UF_COD = {'PR': '41', 'SP': '35', 'SC': '42', 'RS': '43'}


def dv(ch):
    soma, p = 0, 2
    for c in reversed(ch):
        soma += int(c) * p
        p = 2 if p == 9 else p + 1
    d = 11 - soma % 11
    return '0' if d >= 10 else str(d)


def chave(uf, aamm, cnpj, mod, serie, n):
    base = UF_COD[uf] + aamm + cnpj + mod + '%03d' % serie + '%09d' % n + '1' + '%08d' % (n * 7 % 99999999)
    return base + dv(base)


def item(i, cfop, vprod, cst, aliq=0, st=0.0, pis_cst='01', ipi=0.0, orig='0', csosn_cred=0.0, xprod='PRODUTO', ncm='84716052'):
    if cst in ('101', '102', '500', '201'):
        icms = '<ICMSSN%s><orig>%s</orig><CSOSN>%s</CSOSN>%s</ICMSSN%s>' % (
            cst, orig, cst, ('<pCredSN>2.50</pCredSN><vCredICMSSN>%.2f</vCredICMSSN>' % csosn_cred) if cst == '101' else '', cst)
    else:
        vicms = round(vprod * aliq / 100, 2)
        bc = vprod if aliq else 0
        icms = '<ICMS%s><orig>%s</orig><CST>%s</CST><modBC>3</modBC><vBC>%.2f</vBC><pICMS>%.2f</pICMS><vICMS>%.2f</vICMS>%s</ICMS%s>' % (
            cst if cst != '60' else '60', orig, cst, bc, aliq, vicms,
            ('<modBCST>4</modBCST><vBCST>%.2f</vBCST><pICMSST>18.00</pICMSST><vICMSST>%.2f</vICMSST>' % (vprod * 1.4, st)) if st else '', cst)
        if cst == '60':
            icms = '<ICMS60><orig>%s</orig><CST>60</CST><vBCSTRet>0.00</vBCSTRet><vICMSSTRet>0.00</vICMSSTRet></ICMS60>' % orig
    pis_v = round(vprod * 0.0165, 2) if pis_cst in ('01',) else 0
    cof_v = round(vprod * 0.076, 2) if pis_cst in ('01',) else 0
    if pis_cst == '01':
        pis = '<PISAliq><CST>01</CST><vBC>%.2f</vBC><pPIS>1.65</pPIS><vPIS>%.2f</vPIS></PISAliq>' % (vprod, pis_v)
        cof = '<COFINSAliq><CST>01</CST><vBC>%.2f</vBC><pCOFINS>7.60</pCOFINS><vCOFINS>%.2f</vCOFINS></COFINSAliq>' % (vprod, cof_v)
    else:
        pis = '<PISNT><CST>%s</CST></PISNT>' % pis_cst
        cof = '<COFINSNT><CST>%s</CST></COFINSNT>' % pis_cst
    ipi_x = ('<IPI><cEnq>999</cEnq><IPITrib><CST>50</CST><vBC>%.2f</vBC><pIPI>10.00</pIPI><vIPI>%.2f</vIPI></IPITrib></IPI>' % (vprod, ipi)) if ipi else ''
    return ('<det nItem="%d"><prod><cProd>%03d</cProd><cEAN>SEM GTIN</cEAN><xProd>%s</xProd><NCM>%s</NCM><CFOP>%s</CFOP><uCom>UN</uCom><qCom>1.0000</qCom>'
            '<vUnCom>%.2f</vUnCom><vProd>%.2f</vProd><indTot>1</indTot></prod><imposto><ICMS>%s</ICMS>%s<PIS>%s</PIS><COFINS>%s</COFINS></imposto></det>') % (
        i, i, xprod, ncm, cfop, vprod, vprod, icms, ipi_x, pis, cof), {
        'vProd': vprod, 'vICMS': 0 if cst in ('101', '102', '500', '201', '60') else round(vprod * aliq / 100, 2), 'vST': st, 'vIPI': ipi,
        'vPIS': pis_v, 'vCOFINS': cof_v, 'vBC': vprod if aliq and cst not in ('101', '102', '500', '201', '60') else 0}


def nfe(pasta, n, emit, dest, itens, data='2026-09-10', tpNF='1', crt='3', serie=1, fin='1', refs=()):
    ch = chave(emit[2], data[2:4] + data[5:7], emit[0], '55', serie, n)
    dets = ''.join(i[0] for i in itens)
    t = {k: round(sum(i[1][k] for i in itens), 2) for k in itens[0][1]}
    vnf = round(t['vProd'] + t['vST'] + t['vIPI'], 2)
    dest_doc = '<CPF>%s</CPF>' % dest[0] if len(dest[0]) == 11 else '<CNPJ>%s</CNPJ>' % dest[0]
    xml = ('<?xml version="1.0" encoding="UTF-8"?><nfeProc xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00"><NFe><infNFe Id="NFe%s" versao="4.00">'
           '<ide><cUF>%s</cUF><natOp>VENDA</natOp><mod>55</mod><serie>%d</serie><nNF>%d</nNF><dhEmi>%sT10:00:00-03:00</dhEmi><tpNF>%s</tpNF><idDest>%s</idDest><finNFe>%s</finNFe><indFinal>0</indFinal>%s</ide>'
           '<emit><CNPJ>%s</CNPJ><xNome>%s</xNome><enderEmit><cMun>4106902</cMun><UF>%s</UF></enderEmit><IE>123</IE><CRT>%s</CRT></emit>'
           '<dest>%s<xNome>%s</xNome><enderDest><cMun>4106902</cMun><UF>%s</UF></enderDest><indIEDest>1</indIEDest></dest>%s'
           '<total><ICMSTot><vBC>%.2f</vBC><vICMS>%.2f</vICMS><vICMSDeson>0.00</vICMSDeson><vFCP>0.00</vFCP><vBCST>0.00</vBCST><vST>%.2f</vST><vFCPST>0.00</vFCPST>'
           '<vProd>%.2f</vProd><vFrete>0.00</vFrete><vSeg>0.00</vSeg><vDesc>0.00</vDesc><vII>0.00</vII><vIPI>%.2f</vIPI><vIPIDevol>0.00</vIPIDevol><vPIS>%.2f</vPIS><vCOFINS>%.2f</vCOFINS><vOutro>0.00</vOutro><vNF>%.2f</vNF></ICMSTot></total>'
           '</infNFe></NFe><protNFe versao="4.00"><infProt><chNFe>%s</chNFe><cStat>100</cStat></infProt></protNFe></nfeProc>') % (
        ch, UF_COD[emit[2]], serie, n, data, tpNF, '1' if emit[2] == dest[2] else '2', fin, ''.join('<NFref><refNFe>%s</refNFe></NFref>' % r for r in refs), emit[0], emit[1], emit[2], crt, dest_doc, dest[1], dest[2], dets,
        t['vBC'], t['vICMS'], t['vST'], t['vProd'], t['vIPI'], t['vPIS'], t['vCOFINS'], vnf, ch)
    os.makedirs(pasta, exist_ok=True)
    open(os.path.join(pasta, 'NFe%s.xml' % ch), 'w').write(xml)
    return ch, vnf, t


def cancelamento(pasta, ch):
    xml = ('<?xml version="1.0"?><procEventoNFe xmlns="http://www.portalfiscal.inf.br/nfe"><evento><infEvento Id="ID110111%s01"><chNFe>%s</chNFe>'
           '<dhEvento>2026-09-12T10:00:00-03:00</dhEvento><tpEvento>110111</tpEvento><nSeqEvento>1</nSeqEvento><detEvento><descEvento>Cancelamento</descEvento></detEvento>'
           '</infEvento></evento><retEvento><infEvento><cStat>135</cStat></infEvento></retEvento></procEventoNFe>') % (ch, ch)
    open(os.path.join(pasta, 'Canc%s.xml' % ch), 'w').write(xml)


def cte(pasta, n, emit, uf_ini, uf_fim, cfop, valor, aliq, toma, rem_cnpj, nfe_ch=''):
    ch = chave(emit[2], '2609', emit[0], '57', 1, n)
    vicms = round(valor * aliq / 100, 2)
    toma_x = '<toma3><toma>%s</toma></toma3>' % toma
    xml = ('<?xml version="1.0"?><cteProc xmlns="http://www.portalfiscal.inf.br/cte"><CTe><infCte Id="CTe%s"><ide><CFOP>%s</CFOP><natOp>FRETE</natOp><mod>57</mod><serie>1</serie><nCT>%d</nCT>'
           '<dhEmi>2026-09-08T09:00:00-03:00</dhEmi><tpCTe>0</tpCTe><UFIni>%s</UFIni><UFFim>%s</UFFim>%s</ide>'
           '<emit><CNPJ>%s</CNPJ><xNome>%s</xNome><enderEmit><UF>%s</UF></enderEmit></emit>'
           '<rem><CNPJ>%s</CNPJ><xNome>REMETENTE</xNome><enderReme><UF>%s</UF></enderReme></rem><dest><CNPJ>%s</CNPJ><xNome>%s</xNome><enderDest><UF>PR</UF></enderDest></dest>'
           '<vPrest><vTPrest>%.2f</vTPrest><vRec>%.2f</vRec></vPrest><imp><ICMS><ICMS00><CST>00</CST><vBC>%.2f</vBC><pICMS>%.2f</pICMS><vICMS>%.2f</vICMS></ICMS00></ICMS></imp>'
           '<infCTeNorm><infCarga><vCarga>10000.00</vCarga></infCarga><infDoc><infNFe><chave>%s</chave></infNFe></infDoc></infCTeNorm>'
           '</infCte></CTe><protCTe><infProt><cStat>100</cStat></infProt></protCTe></cteProc>') % (
        ch, cfop, n, uf_ini, uf_fim, toma_x, emit[0], emit[1], emit[2], rem_cnpj, uf_ini, EMP['cnpj'], EMP['nome'], valor, valor, valor, aliq, vicms, nfe_ch)
    os.makedirs(pasta, exist_ok=True)
    open(os.path.join(pasta, 'CTe%s.xml' % ch), 'w').write(xml)
    return ch


def nfse_nac(pasta, n, prest, toma, valor, aliq, retido=False, irrf=0.0):
    viss = round(valor * aliq / 100, 2)
    xml = ('<?xml version="1.0"?><NFSe xmlns="http://www.sped.fazenda.gov.br/nfse"><infNFSe Id="NFS4106902%s%013d"><nNFSe>%d</nNFSe><dhProc>2026-09-15T10:00:00-03:00</dhProc>'
           '<emit><CNPJ>%s</CNPJ><xNome>%s</xNome></emit><valores><vBC>%.2f</vBC><pAliqAplic>%.2f</pAliqAplic><vISSQN>%.2f</vISSQN><vLiq>%.2f</vLiq></valores>'
           '<DPS><infDPS><dhEmi>2026-09-15T10:00:00-03:00</dhEmi><dCompet>2026-09-15</dCompet><prest><CNPJ>%s</CNPJ></prest><toma><CNPJ>%s</CNPJ><xNome>%s</xNome></toma>'
           '<serv><cServ><cTribNac>170101</cTribNac><xDescServ>CONSULTORIA</xDescServ></cServ></serv><valores><vServPrest><vServ>%.2f</vServ></vServPrest>'
           '<trib><tribMun><tribISSQN>1</tribISSQN><pAliq>%.2f</pAliq><tpRetISSQN>%s</tpRetISSQN></tribMun><tribFed><vRetIRRF>%.2f</vRetIRRF></tribFed></trib></valores></infDPS></DPS></infNFSe></NFSe>') % (
        prest[0], n, n, prest[0], prest[1], valor, aliq, viss, valor - irrf - (viss if retido else 0), prest[0], toma[0], toma[1], valor, aliq, '2' if retido else '1', irrf)
    os.makedirs(pasta, exist_ok=True)
    open(os.path.join(pasta, 'NFSe_%s_%d.xml' % (prest[0], n)), 'w').write(xml)


def lote_abrasf(pasta, notas):
    comps = ''
    for n, prest, valor, aliq, retido, ir in notas:
        viss = round(valor * aliq / 100, 2)
        comps += ('<CompNfse><Nfse><InfNfse><Numero>%d</Numero><CodigoVerificacao>AB%dCD</CodigoVerificacao><DataEmissao>2026-09-20T08:00:00</DataEmissao><Competencia>2026-09-01</Competencia>'
                  '<Servico><Valores><ValorServicos>%.2f</ValorServicos><IssRetido>%s</IssRetido><ValorIss>%.2f</ValorIss><Aliquota>%.2f</Aliquota><ValorIr>%.2f</ValorIr><ValorLiquidoNfse>%.2f</ValorLiquidoNfse></Valores>'
                  '<ItemListaServico>7.10</ItemListaServico><Discriminacao>MANUTENCAO PREDIAL</Discriminacao></Servico>'
                  '<PrestadorServico><IdentificacaoPrestador><CpfCnpj><Cnpj>%s</Cnpj></CpfCnpj></IdentificacaoPrestador><RazaoSocial>%s</RazaoSocial></PrestadorServico>'
                  '<TomadorServico><IdentificacaoTomador><CpfCnpj><Cnpj>%s</Cnpj></CpfCnpj></IdentificacaoTomador><RazaoSocial>%s</RazaoSocial></TomadorServico>'
                  '</InfNfse></Nfse></CompNfse>') % (n, n, valor, '1' if retido else '2', viss, aliq, ir, valor - ir, prest[0], prest[1], EMP['cnpj'], EMP['nome'])
    os.makedirs(pasta, exist_ok=True)
    open(os.path.join(pasta, 'lote_tomados_abrasf.xml'), 'w').write('<?xml version="1.0"?><ConsultarNfseResposta><ListaNfse>%s</ListaNfse></ConsultarNfseResposta>' % comps)


def danfe_pdf(caminho, paginas):
    import fitz
    doc = fitz.open()
    for titulo, ch, extra in paginas:
        pg = doc.new_page()
        txt = '%s\n\nCHAVE DE ACESSO\n%s\n\n%s' % (titulo, ' '.join(ch[i:i + 4] for i in range(0, 44, 4)), extra)
        pg.insert_text((50, 80), txt, fontsize=11)
    doc.save(caminho)


def main():
    if os.path.isdir(BASE):
        shutil.rmtree(BASE)
    emp = (EMP['cnpj'], EMP['nome'], EMP['uf'])
    E, S = os.path.join(BASE, 'entrada'), os.path.join(BASE, 'saida')
    L_ent, L_sai, L_frete, L_prest, L_toma = [], [], [], [], []

    def lanc(lst, n, data_emi, data_ent, part, linhas, serie='1'):
        for k, l in enumerate(linhas):
            lst.append([n if k == 0 else '', serie if k == 0 else '', data_emi if k == 0 else '', data_ent if k == 0 else '',
                        part[0] if k == 0 else '', part[1] if k == 0 else '', part[2] if k == 0 else ''] + l)

    def L(cfop, cst, vc, bc, aliq, icms, ipi=0, st=0, cst_pis='50', base_pis=None, pis=None, cof=None):
        if base_pis is None:
            base_pis = round(vc - icms, 2) if cst_pis == '50' else 0
        if pis is None:
            pis = round(base_pis * 0.0165, 2) if cst_pis == '50' else 0
        if cof is None:
            cof = round(base_pis * 0.076, 2) if cst_pis == '50' else 0
        return [cfop, cst, vc, bc, aliq, icms, ipi, st, cst_pis, base_pis, pis, cof]

    f = lambda k: F[k][:3]
    # E1 correta
    nfe(E, 1001, f('F1'), emp, [item(1, '5102', 1000, '00', 18)])
    lanc(L_ent, 1001, '10/09/2026', '11/09/2026', F['F1'], [L('1102', '000', 1000, 1000, 18, 180)])
    # E2 interestadual lançada como estadual
    nfe(E, 2002, f('F2'), emp, [item(1, '6102', 2000, '00', 12)])
    lanc(L_ent, 2002, '10/09/2026', '12/09/2026', F['F2'], [L('1102', '000', 2000, 2000, 12, 240)])
    # E3 compra com ST lançada em 1102
    nfe(E, 3003, f('F5'), emp, [item(1, '5403', 800, '10', 18, st=57.6, xprod='PASTILHA DE FREIO')])
    lanc(L_ent, 3003, '10/09/2026', '12/09/2026', F['F5'], [L('1102', '010', 857.6, 800, 18, 144)])
    # E4 fornecedor do Simples (CSOSN 101) lançado com CST 000 e crédito cheio
    nfe(E, 4004, f('F3'), emp, [item(1, '5102', 1000, '101', csosn_cred=25.0)], crt='1')
    lanc(L_ent, 4004, '10/09/2026', '13/09/2026', F['F3'], [L('1102', '000', 1000, 1000, 18, 180)])
    # E5 uso e consumo com crédito
    nfe(E, 5005, f('F1'), emp, [item(1, '5102', 200, '00', 18, xprod='PAPEL A4')])
    lanc(L_ent, 5005, '10/09/2026', '13/09/2026', F['F1'], [L('1556', '000', 200, 200, 18, 36, cst_pis='70')])
    # E6 monofásico com crédito de PIS/COFINS
    nfe(E, 6006, f('F4'), emp, [item(1, '6102', 3000, '00', 12, pis_cst='04', xprod='CERVEJA LATA 350ML', ncm='22030000')])
    lanc(L_ent, 6006, '10/09/2026', '14/09/2026', F['F4'], [L('2102', '000', 3000, 3000, 12, 360)])
    # E7 duas linhas na nota, só uma lançada
    nfe(E, 7007, f('F1'), emp, [item(1, '5102', 1200, '00', 18, xprod='MONITOR 24'), item(2, '5405', 300, '60', xprod='CABO HDMI')])
    lanc(L_ent, 7007, '10/09/2026', '14/09/2026', F['F1'], [L('1102', '000', 1200, 1200, 18, 216)])
    # E8 não lançada
    nfe(E, 8008, f('F2'), emp, [item(1, '6102', 4500, '00', 12)])
    # E9 cancelada e lançada
    ch9, _, _ = nfe(E, 9009, f('F1'), emp, [item(1, '5102', 700, '00', 18)])
    cancelamento(E, ch9)
    lanc(L_ent, 9009, '10/09/2026', '15/09/2026', F['F1'], [L('1102', '000', 700, 700, 18, 126)])
    # E10 lançada em duplicidade (blocos separados)
    nfe(E, 1010, f('F1'), emp, [item(1, '5102', 650, '00', 18)])
    lanc(L_ent, 1010, '10/09/2026', '15/09/2026', F['F1'], [L('1102', '000', 650, 650, 18, 117)])
    # E11 valor divergente
    nfe(E, 1111, f('F1'), emp, [item(1, '5102', 1500, '00', 18)])
    lanc(L_ent, 1111, '10/09/2026', '16/09/2026', F['F1'], [L('1102', '000', 1550, 1500, 18, 270)])
    # E12 base PIS com ICMS e data de entrada antes da emissão
    nfe(E, 1212, f('F1'), emp, [item(1, '5102', 5000, '00', 18)], data='2026-09-18')
    lanc(L_ent, 1212, '18/09/2026', '17/09/2026', F['F1'], [L('1102', '000', 5000, 5000, 18, 900, base_pis=5000)])
    # posto de SP emite 5656 (venda presencial); 2653 está certo
    nfe(E, 1313, f('F6'), emp, [item(1, '5656', 250, '60', xprod='GASOLINA COMUM', ncm='27101259', pis_cst='04')])
    lanc(L_ent, 1313, '10/09/2026', '11/09/2026', F['F6'], [L('2653', '060', 250, 0, 0, 0, cst_pis='70')])
    # mesmo posto lançado com 1653: conferir
    nfe(E, 1314, f('F6'), emp, [item(1, '5656', 180, '60', xprod='ETANOL HIDRATADO', ncm='22071090', pis_cst='04')])
    lanc(L_ent, 1314, '10/09/2026', '11/09/2026', F['F6'], [L('1653', '060', 180, 0, 0, 0, cst_pis='70')])
    # refeição em SP (5102) lançada 2556: certo
    nfe(E, 1315, f('F8'), emp, [item(1, '5102', 95, '00', 0, xprod='REFEICAO VENDA A VISTA', ncm='21069090', pis_cst='01')])
    lanc(L_ent, 1315, '10/09/2026', '11/09/2026', F['F8'], [L('2556', '000', 95, 0, 0, 0, cst_pis='70')])
    # uso e consumo com crédito de PIS/COFINS e IPI
    nfe(E, 1316, f('F1'), emp, [item(1, '5102', 400, '00', 18, ipi=40, xprod='DETERGENTE NEUTRO 5L', ncm='34022000')])
    lanc(L_ent, 1316, '10/09/2026', '11/09/2026', F['F1'], [L('1556', '000', 440, 0, 0, 0, ipi=40, cst_pis='50', base_pis=440)])
    # empilhadeira lançada como industrialização
    nfe(E, 1317, f('F7'), emp, [item(1, '6101', 85000, '00', 12, xprod='EMPILHADEIRA ELETRICA 2T', ncm='84271011')])
    lanc(L_ent, 1317, '10/09/2026', '11/09/2026', F['F7'], [L('2101', '000', 85000, 85000, 12, 10200)])
    # papel A4 lançado como compra para comercialização
    nfe(E, 1318, f('F1'), emp, [item(1, '5102', 300, '00', 18, xprod='PAPEL A4 SULFITE 75G', ncm='48025610')])
    lanc(L_ent, 1318, '10/09/2026', '11/09/2026', F['F1'], [L('1102', '000', 300, 300, 18, 54)])
    # CST 60 (ST) lançado em 1556: tem que ser 1407
    nfe(E, 1319, f('F5'), emp, [item(1, '5405', 120, '60', xprod='LAMPADA LED 9W', ncm='85395200')])
    lanc(L_ent, 1319, '10/09/2026', '11/09/2026', F['F5'], [L('1556', '060', 120, 0, 0, 0, cst_pis='70')])
    # sem ST lançado em 1407: tem que ser 1556
    nfe(E, 1320, f('F1'), emp, [item(1, '5102', 90, '00', 0, xprod='CAFE TORRADO 500G', ncm='09012100')])
    lanc(L_ent, 1320, '10/09/2026', '11/09/2026', F['F1'], [L('1407', '000', 90, 0, 0, 0, cst_pis='70')])
    # nota mista (com e sem ST) toda numa linha só 1556
    nfe(E, 1321, f('F5'), emp, [item(1, '5102', 200, '00', 0, xprod='PANO DE CHAO'), item(2, '5405', 80, '60', xprod='PILHA AA')])
    lanc(L_ent, 1321, '10/09/2026', '11/09/2026', F['F5'], [L('1556', '000', 280, 0, 0, 0, cst_pis='70')])
    # nota mista certa: duas linhas
    nfe(E, 1322, f('F5'), emp, [item(1, '5102', 200, '00', 0, xprod='PANO DE CHAO'), item(2, '5405', 80, '60', xprod='PILHA AA')])
    lanc(L_ent, 1322, '10/09/2026', '11/09/2026', F['F5'], [L('1556', '000', 200, 0, 0, 0, cst_pis='70'), L('1407', '060', 80, 0, 0, 0, cst_pis='70')])
    # duplicidade de E10 mais abaixo
    lanc(L_ent, 1010, '10/09/2026', '20/09/2026', F['F1'], [L('1102', '000', 650, 650, 18, 117)])
    # lançamento sem XML, mas com DANFE no PDF
    ch_pdf1 = chave('PR', '2609', F['F1'][0], '55', 1, 9999)
    lanc(L_ent, 9999, '05/09/2026', '06/09/2026', F['F1'], [L('1102', '000', 450, 450, 18, 81)])
    # lançamento sem XML e sem PDF
    lanc(L_ent, 8888, '05/09/2026', '06/09/2026', F['F2'], [L('2102', '000', 980, 980, 12, 117.6)])
    # DANFE no PDF sem lançamento
    ch_pdf2 = chave('PR', '2609', F['F5'][0], '55', 1, 7777)
    ch1 = chave('PR', '2609', F['F1'][0], '55', 1, 1001)

    # Saídas
    ch_s1, _, _ = nfe(S, 501, emp, C['C1'], [item(1, '5102', 2500, '00', 18)])
    lanc(L_sai, 501, '10/09/2026', '10/09/2026', C['C1'], [L('5102', '000', 2500, 2500, 18, 450, cst_pis='01', base_pis=2500, pis=41.25, cof=190)])
    nfe(S, 502, emp, C['C2'], [item(1, '6102', 3000, '00', 12)])
    lanc(L_sai, 502, '10/09/2026', '10/09/2026', C['C2'], [L('5102', '000', 3000, 3000, 12, 360, cst_pis='01', base_pis=3000, pis=49.5, cof=228)])
    nfe(S, 503, emp, C['C1'], [item(1, '5405', 900, '60', xprod='FILTRO DE OLEO')])
    lanc(L_sai, 503, '10/09/2026', '10/09/2026', C['C1'], [L('5405', '000', 900, 0, 0, 0, cst_pis='01', base_pis=900, pis=14.85, cof=68.4)])
    nfe(S, 504, emp, C['CPF'], [item(1, '5102', 1000, '00', 18)])
    lanc(L_sai, 504, '10/09/2026', '10/09/2026', C['CPF'], [L('5102', '000', 1000, 1000, 18, 0, cst_pis='01', base_pis=1000, pis=16.5, cof=76)])
    nfe(S, 505, emp, C['C1'], [item(1, '5102', 1800, '00', 18)])
    # numeração: 510 emitida, 506-507 inutilizadas, 508-509 pulados
    nfe(S, 510, emp, C['C1'], [item(1, '5102', 400, '00', 18)])
    lanc(L_sai, 510, '20/09/2026', '20/09/2026', C['C1'], [L('5102', '000', 400, 400, 18, 72, cst_pis='01', base_pis=400, pis=6.6, cof=30.4)])
    open(os.path.join(S, 'inut_506_507.xml'), 'w').write('<?xml version="1.0"?><procInutNFe xmlns="http://www.portalfiscal.inf.br/nfe"><inutNFe><infInut Id="ID41"><tpAmb>1</tpAmb><xServ>INUTILIZAR</xServ><cUF>41</cUF><ano>26</ano><CNPJ>%s</CNPJ><mod>55</mod><serie>1</serie><nNFIni>506</nNFIni><nNFFin>507</nNFFin><xJust>FALHA NO SISTEMA EMISSOR</xJust></infInut></inutNFe><retInutNFe><infInut><cStat>102</cStat><dhRecbto>2026-09-21T10:00:00-03:00</dhRecbto></infInut></retInutNFe></procInutNFe>' % EMP['cnpj'])
    lanc(L_sai, 506, '21/09/2026', '21/09/2026', C['C1'], [L('5102', '000', 350, 350, 18, 63, cst_pis='01', base_pis=350, pis=5.78, cof=26.6)])
    # devolução de venda própria (entrada 1202) da 501: devolve 2 de 5 com ICMS errado
    nfe(S, 511, emp, C['C1'], [item(1, '1202', 1000, '00', 12)], tpNF='0', fin='4', refs=[ch_s1], data='2026-09-22')
    lanc(L_ent, 511, '22/09/2026', '22/09/2026', C['C1'], [L('1202', '000', 1000, 1000, 12, 120)])
    # devolução de venda emitida pelo cliente (5202) da 502, lançada em 2102
    ch_s2 = chave('PR', '2609', EMP['cnpj'], '55', 1, 502)
    nfe(E, 777, (C['C2'][0], C['C2'][1], 'RS'), emp, [item(1, '5202', 3000, '00', 12)], fin='4', refs=[ch_s2], data='2026-09-23')
    lanc(L_ent, 777, '23/09/2026', '24/09/2026', C['C2'], [L('2102', '000', 3000, 3000, 12, 360)])
    # devolução de compra própria (5202) da E1 1001 com preço diferente
    nfe(S, 512, emp, F['F1'][:3], [item(1, '5202', 1100, '00', 18)], fin='4', refs=[ch1], data='2026-09-24')
    lanc(L_sai, 512, '24/09/2026', '24/09/2026', F['F1'], [L('5202', '000', 1100, 1100, 18, 198, cst_pis='50', base_pis=902, pis=14.88, cof=68.55)])
    nfe(S, 499, emp, C['C1'], [item(1, '5102', 600, '00', 18)], data='2026-08-30')

    # Frete
    FR = os.path.join(BASE, 'frete')
    cte(FR, 301, T['T1'], 'PR', 'PR', '5353', 500, 12, '3', F['F1'][0], ch1)
    L_frete.append([301, '1', '08/09/2026', '09/09/2026', T['T1'][0], T['T1'][1], 'PR', '1353', '000', 500, 500, 12, 60, 0, 0, '50', 500, 8.25, 38])
    cte(FR, 302, T['T2'], 'SP', 'PR', '6353', 800, 12, '3', F['F2'][0])
    L_frete.append([302, '1', '08/09/2026', '09/09/2026', T['T2'][0], T['T2'][1], 'SP', '1353', '000', 800, 0, 0, 0, 0, 0, '70', 0, 0, 0])
    cte(FR, 303, T['T1'], 'PR', 'PR', '5353', 300, 12, '0', F['F1'][0])
    L_frete.append([303, '1', '08/09/2026', '09/09/2026', T['T1'][0], T['T1'][1], 'PR', '1353', '000', 300, 300, 12, 36, 0, 0, '50', 300, 4.95, 22.8])

    # Prestados
    P = os.path.join(BASE, 'prestado')
    nfse_nac(P, 101, emp, C['C1'], 10000, 5)
    L_prest.append([101, '15/09/2026', C['C1'][0], C['C1'][1], 10000, 5, 500, 'N', 0, 0, 0, 0, 0])
    nfse_nac(P, 102, emp, C['C2'], 8000, 5, retido=True, irrf=120)
    L_prest.append([102, '15/09/2026', C['C2'][0], C['C2'][1], 8000, 5, 400, 'N', 120, 0, 0, 0, 0])
    nfse_nac(P, 103, emp, C['C1'], 3300, 5)
    L_prest.append([103, '15/09/2026', C['C1'][0], C['C1'][1], 3000, 5, 150, 'N', 0, 0, 0, 0, 0])

    # Tomados
    TO = os.path.join(BASE, 'tomado')
    pr1 = ('12312312000155', 'MANUTENCAO PREDIAL XYZ LTDA', 'PR')
    pr2 = ('32132132000166', 'LIMPEZA E CONSERVACAO LTDA', 'PR')
    pr3 = ('45645645000177', 'CONTABILIDADE MODELO LTDA', 'PR')
    lote_abrasf(TO, [(55, pr1, 2000, 5, True, 0), (56, pr2, 2000, 3, False, 30)])
    L_toma.append([55, '20/09/2026', pr1[0], pr1[1], 2000, 5, 100, 'N', 0, 0, 0, 0, 0])
    L_toma.append([56, '20/09/2026', pr2[0], pr2[1], 2000, 3, 60, 'N', 0, 0, 0, 0, 0])
    nfse_nac(TO, 900, pr3, emp, 1500, 2)
    L_toma.append([900, '15/09/2026', pr3[0], pr3[1], 1500, 2, 30, 'N', 0, 0, 0, 0, 0])

    # PDF diversos
    D = os.path.join(BASE, 'diversos'); os.makedirs(D, exist_ok=True)
    danfe_pdf(os.path.join(D, 'Notas_diversas_setembro.pdf'), [
        ('DANFE - DOCUMENTO AUXILIAR DA NOTA FISCAL ELETRONICA\nDISTRIBUIDORA ALFA LTDA', ch_pdf1, 'NF-e Nº 9.999  Série 1   VALOR TOTAL DA NOTA 450,00'),
        ('DANFE - DOCUMENTO AUXILIAR DA NOTA FISCAL ELETRONICA\nEPSILON AUTOPECAS LTDA', ch_pdf2, 'NF-e Nº 7.777  Série 1   VALOR TOTAL DA NOTA 1.230,00'),
        ('DANFE - DOCUMENTO AUXILIAR DA NOTA FISCAL ELETRONICA\nDISTRIBUIDORA ALFA LTDA', ch1, 'NF-e Nº 1.001'),
    ])

    # Planilha do sistema (layout de "livro" com título e linhas de continuação)
    PL = os.path.join(BASE, 'planilha'); os.makedirs(PL, exist_ok=True)
    wb = openpyxl.Workbook()
    cab_nf = ['Nº Nota', 'Série', 'Dt. Emissão', 'Dt. Entrada', 'CNPJ/CPF', 'Fornecedor', 'UF', 'CFOP', 'CST', 'Valor Contábil', 'Base ICMS', 'Alíq.',
              'Valor ICMS', 'Valor IPI', 'ICMS ST', 'CST PIS', 'Base PIS/COFINS', 'Valor PIS', 'Valor COFINS']
    ws = wb.active; ws.title = 'Entradas'
    ws.append(['COMERCIAL EXEMPLO LTDA - Relação de notas de entrada - 09/2026']); ws.append([])
    ws.append(cab_nf)
    for r in L_ent:
        ws.append(r)
    ws2 = wb.create_sheet('Saídas')
    ws2.append(['Relação de notas de saída - 09/2026']); ws2.append([])
    ws2.append([c.replace('Dt. Entrada', 'Dt. Saída').replace('Fornecedor', 'Cliente') for c in cab_nf])
    for r in L_sai:
        ws2.append(r)
    ws3 = wb.create_sheet('Fretes CT-e')
    ws3.append(['Nº Nota', 'Série', 'Dt. Emissão', 'Dt. Entrada', 'CNPJ/CPF', 'Transportadora', 'UF', 'CFOP', 'CST', 'Valor Contábil', 'Base ICMS', 'Alíq.',
                'Valor ICMS', 'Valor IPI', 'ICMS ST', 'CST PIS', 'Base PIS/COFINS', 'Valor PIS', 'Valor COFINS'])
    for r in L_frete:
        ws3.append(r)
    wb.save(os.path.join(PL, 'planilha_sistema_09-2026.xlsx'))
    # relatório de serviços (uma aba, coluna Tipo prestado/tomado, PIS/COFINS = retidos)
    PS = os.path.join(BASE, 'planilha_servicos'); os.makedirs(PS, exist_ok=True)
    wb2 = openpyxl.Workbook(); w = wb2.active; w.title = 'Servicos'
    w.append(['Relatório de Serviços - 09/2026']); w.append([])
    w.append(['Tipo', 'Nº NFS-e', 'Data', 'CNPJ/CPF', 'Prestador/Tomador', 'Cód. Serviço', 'Valor Serviços', 'Base ISS', 'Alíq. ISS', 'Valor ISS', 'ISS Retido', 'IRRF', 'CSLL', 'PIS', 'COFINS', 'INSS'])
    cod = {101: '17.01', 102: '17.01', 103: '17.02', 55: '7.10', 56: '7.10', 900: '17.19'}
    for tipo, lst in (('Prestado', L_prest), ('Tomado', L_toma)):
        for r in lst:
            n, data, cnpj, nome, vs, aliq, iss, ret, ir, csll, pis, cof, inss = r
            w.append([tipo, n, data, cnpj, nome, cod.get(n, ''), vs, vs, aliq, iss, ret, ir, csll, pis, cof, inss])
    wb2.save(os.path.join(PS, 'relatorio_servicos_09-2026.xlsx'))
    print('lote gerado em', BASE)


if __name__ == '__main__':
    main()
